import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from io import BytesIO
from zipfile import ZipFile
import pytest
from PIL import Image
from docx import Document
from pypdf import PdfReader
from model import *
from exports import make_docx,make_pdf

def example(kind='plan'):
    d=new_document();d.update(kind=kind,program='학술 교류 프로그램',event_name='학술세미나',subject='임상 경험과 연구 공유',places=['세미나실'],people=['교수진','대학원생'],speakers='초청 교수',purpose='학술 교류를 통해 연구 역량을 강화한다.',details='초청 강연 및 증례 토론\n참석자 질의응답',expected='공동 연구 활성화 및 전문성 향상',summary='초청 강연과 증례 토론을 진행하였다.\n향후 공동 연구 방향을 논의하였다.',budget=[dict(item='연자료',amount=500000,note='교육비'),dict(item='운영비',amount=200000,note='행사 운영')])
    raw=BytesIO();Image.new('RGB',(1200,600),'#137f75').save(raw,'PNG');p=photo_record(raw.getvalue(),'photo.png');p['caption']='행사 사진 예시';d['photos']=[p];return d

def test_copy_and_backup():
    d=example();c=copy_document(d);c['photos'][0]['caption']='수정';assert d['photos'][0]['caption']!='수정';assert c['id']!=d['id']
    phrases=[dict(id='p1',category='purpose',name='목적',content='교류 활성화')]
    restored=read_backup(make_backup(d,[d],phrases));assert restored['draft']==d;assert restored['phrases']==phrases
    with pytest.raises(ValueError):read_backup(make_backup(d,[d,d],phrases))

def test_image_validation():
    raw=BytesIO();Image.new('RGBA',(3000,1000),(0,0,0,0)).save(raw,'PNG')
    result=Image.open(BytesIO(normalize_image(raw.getvalue())));assert result.size==(2000,667);assert result.getpixel((0,0))==(255,255,255)
    with pytest.raises(Exception):normalize_image(b'not an image')
    d=example();d['include_time']=True;d['start']='12:00';d['end']='10:00'
    with pytest.raises(ValueError):validate_document(d)

@pytest.mark.parametrize('kind',['plan','report'])
def test_exports(kind):
    d=example(kind);validate_document(d)
    doc=Document(BytesIO(make_docx(d)));text='\n'.join(p.text for p in doc.paragraphs)
    assert '{{' not in text;assert '행사 사진 예시' in text;assert len(doc.inline_shapes)==1
    if kind=='plan':assert doc.tables[0].cell(3,1).text=='700,000원'
    pdf=PdfReader(BytesIO(make_pdf(d)));text='\n'.join(p.extract_text() for p in pdf.pages)
    assert '학술세미나' in text;assert '행사 사진 예시' in text
    d['details']=('긴 문장으로 페이지 나눔을 확인합니다. '*15+'\n')*30
    if kind=='report':d['summary']=d['details']
    assert len(PdfReader(BytesIO(make_pdf(d))).pages)>2

def test_streamlit_wizard():
    from streamlit.testing.v1 import AppTest
    at=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=30).run()
    assert not at.exception
    d=example();at.session_state['draft']=d;at.run()
    def click(label):
        next(b for b in at.button if b.label==label).click().run();assert not at.exception
    click('다음 →');click('다음 →')
    at.text_area(key='input_purpose').set_value('상용구 테스트 목적').run()
    at.text_input(key='phrase_name_purpose').set_value('기본 목적').run()
    at.button(key='save_phrase_purpose').click().run();assert len(at.session_state['phrases'])==1
    click('다음 →');click('다음 →');click('현재 세션에 문서 보관');click('PDF·Word 파일 준비')
    assert set(at.session_state['export']['files'])=={'pdf','docx'}
    at.radio(key='navigation').set_value('상용구 관리').run();assert not at.exception
    at.radio(key='navigation').set_value('새 문서').run();assert not at.exception
    click('← 이전');click('← 이전');assert at.text_area(key='input_purpose').value=='상용구 테스트 목적'

def test_report_wizard_pending_text():
    from streamlit.testing.v1 import AppTest
    at=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=30)
    at.session_state['draft']=example('report');at.run()
    def click(label):
        next(b for b in at.button if b.label==label).click().run();assert not at.exception
    click('다음 →');click('다음 →')
    at.text_area(key='input_summary').set_value('마지막 입력도 저장되어야 합니다.')
    click('다음 →');click('다음 →');click('PDF·Word 파일 준비')
    assert at.session_state['draft']['summary']=='마지막 입력도 저장되어야 합니다.'
    assert set(at.session_state['export']['files'])=={'pdf','docx'}
